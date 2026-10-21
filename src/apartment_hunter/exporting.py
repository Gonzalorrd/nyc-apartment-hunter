import csv
import io
import json
import math
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from apartment_hunter.config import preferences, today
from apartment_hunter.schemas import CONTACT_STATUSES, SCHEMAS
from apartment_hunter.scoring import rank_all
from apartment_hunter.services import rows


def safe_cell(value):
    if isinstance(value, (dict, list)):
        return json.dumps(value, default=str)
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value  # Never execute imported text as spreadsheet formulas.
    return value


def csv_bytes(records: list[dict], headers: list[str] | None = None) -> bytes:
    stream = io.StringIO(newline="")
    fields = headers or (list(records[0]) if records else ["No records"])
    writer = csv.DictWriter(stream, fieldnames=fields, extrasaction="ignore")
    writer.writeheader()
    for row in records:
        writer.writerow({key: safe_cell(row.get(key)) for key in fields})
    return stream.getvalue().encode("utf-8-sig")


def workbook_bytes(conn) -> bytes:
    """Rebuilt snapshot, never an editable second database."""
    book = Workbook()
    book.remove(book.active)
    ranking = rank_all(conn)
    buildings, units = rows(conn, "buildings"), rows(conn, "units")
    summary = [
        {"metric": "Export date", "value": today()},
        {"metric": "Workbook use", "value": "READ-ONLY SNAPSHOT. Edit the app; export again."},
        {"metric": "Buildings tracked", "value": len(buildings)},
        {
            "metric": "Listed units (may need re-verification)",
            "value": sum(u["status"] == "listed" for u in units),
        },
        {"metric": "Eligible units", "value": sum(r["eligibility"] == "eligible" for r in ranking)},
        {
            "metric": "Needs verification",
            "value": sum(r["eligibility"] == "needs_verification" for r in ranking),
        },
        {"metric": "Unknown fees", "value": sum(u["mandatory_fees"] is None for u in units)},
        {"metric": "Budget including mandatory fees (USD/month)", "value": preferences().budget},
        {"metric": "Destination", "value": preferences().destination},
    ]
    configs = [{"setting": k, "value": v} for k, v in preferences().model_dump(mode="json").items()]
    sheets = [
        ("Dashboard", summary, ["metric", "value"]),
        ("Buildings", buildings, list(SCHEMAS["buildings"].model_fields)),
        (
            "Available Units",
            [u for u in units if u["status"] == "listed"],
            list(SCHEMAS["units"].model_fields),
        ),
        (
            "Apartment Rankings",
            ranking,
            ["unit_id", "building", "unit", "eligibility", "score", "completeness_percent"],
        ),
        ("Commutes", rows(conn, "commutes"), list(SCHEMAS["commutes"].model_fields)),
        ("Building Risk Research", rows(conn, "risks"), list(SCHEMAS["risks"].model_fields)),
        ("Leasing Contacts", rows(conn, "contacts"), list(SCHEMAS["contacts"].model_fields)),
        ("Interactions", rows(conn, "interactions"), list(SCHEMAS["interactions"].model_fields)),
        ("Tours", rows(conn, "tours"), list(SCHEMAS["tours"].model_fields)),
        (
            "Follow-up Tasks",
            rows(conn, "tasks")
            + [
                {
                    "building_id": c["building_id"],
                    "due_date": c["next_follow_up"],
                    "summary": f"Contact follow-up #{c['id']}: {c['status']}",
                    "status": c["status"],
                }
                for c in rows(conn, "contacts")
                if c["next_follow_up"] and c["status"] not in ("leased", "archived", "rejected")
            ],
            list(SCHEMAS["tasks"].model_fields),
        ),
        (
            "Price History",
            rows(conn, "price_history"),
            [
                "unit_id",
                "observed_at",
                "base_rent",
                "mandatory_fees",
                "advertised_rent",
                "source_url",
            ],
        ),
        ("Configuration", configs, ["setting", "value"]),
    ]
    for index, (name, records, empty_headers) in enumerate(sheets, 1):
        sheet = book.create_sheet(name)
        fields = (
            list(dict.fromkeys(k for r in records for k in r))
            if records
            else ["id", *empty_headers]
        )
        if name == "Dashboard":
            fields = empty_headers
        sheet.append([name])
        sheet.append(["Export only · USD · Unknown values are blank · Re-export after changes"])
        sheet.append([key.replace("_", " ").title() for key in fields])
        for row in records:
            sheet.append([safe_cell(row.get(key)) for key in fields])
        sheet.freeze_panes = "C4"
        sheet.sheet_view.showGridLines = False
        sheet.row_dimensions[1].height = 30
        sheet.row_dimensions[2].height = 24
        sheet.row_dimensions[3].height = 38
        sheet["A1"].font = Font(name="Calibri", size=20, bold=True, color="173C4A")
        sheet["A2"].font = Font(name="Calibri", size=10, italic=True, color="586B74")
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(2, len(fields)))
        sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(2, len(fields)))
        for col, field in enumerate(fields, 1):
            letter = get_column_letter(col)
            width = (
                42
                if any(
                    t in field
                    for t in ("notes", "summary", "url", "description", "warnings", "reasons")
                )
                else 24
            )
            sheet.column_dimensions[letter].width = width
            header = sheet.cell(3, col)
            header.fill = PatternFill("solid", fgColor="173C4A")
            header.font = Font(name="Calibri", bold=True, color="FFFFFF")
            header.alignment = Alignment(wrap_text=True, vertical="center")
            for rownum in range(4, sheet.max_row + 1):
                cell = sheet.cell(rownum, col)
                cell.font = Font(name="Calibri", size=11, color="172E38")
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(cell.value, (datetime, date)):
                    cell.number_format = "mm/dd/yy"
                elif any(t in field for t in ("rent", "fees", "cost", "deposit")) and isinstance(
                    cell.value, (int, float)
                ):
                    cell.number_format = '"$"#,##0.00'
                elif isinstance(cell.value, float):
                    cell.number_format = "0.0"
                if isinstance(cell.value, str) and cell.value.startswith(("https://", "http://")):
                    cell.hyperlink = cell.value
                    cell.font = Font(name="Calibri", size=11, color="167B99", underline="single")
            if "score" in field and records:
                sheet.conditional_formatting.add(
                    f"{letter}4:{letter}{sheet.max_row}",
                    ColorScaleRule(
                        start_type="num",
                        start_value=0,
                        start_color="FDE4DF",
                        end_type="num",
                        end_value=100,
                        end_color="D7EFE7",
                    ),
                )
            if name == "Leasing Contacts" and field == "status":
                validation = DataValidation(
                    type="list", formula1='"' + ",".join(CONTACT_STATUSES) + '"'
                )
                sheet.add_data_validation(validation)
                validation.add(f"{letter}4:{letter}{max(4, sheet.max_row)}")
        end = f"{get_column_letter(len(fields))}{max(3, sheet.max_row)}"
        sheet.auto_filter.ref = f"A3:{end}"
        if records:
            table = Table(displayName=f"Export{index}", ref=f"A3:{end}")
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            sheet.add_table(table)
        for rownum in range(4, sheet.max_row + 1):
            wrapped_lines = max(
                math.ceil(
                    len(str(sheet.cell(rownum, col).value or ""))
                    / max(8, sheet.column_dimensions[get_column_letter(col)].width - 3)
                )
                for col in range(1, len(fields) + 1)
            )
            sheet.row_dimensions[rownum].height = min(400, max(32, wrapped_lines * 15 + 8))
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_title_rows = "1:3"
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def export_workbook(conn, destination: str | Path) -> Path:
    path = Path(destination)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(workbook_bytes(conn))
    return path
