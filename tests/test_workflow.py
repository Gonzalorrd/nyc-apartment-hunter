import io
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from apartment_hunter.config import today
from apartment_hunter.db import database
from apartment_hunter.demo import load_demo
from apartment_hunter.exporting import csv_bytes, workbook_bytes
from apartment_hunter.providers import refresh_sources
from apartment_hunter.services import (
    import_csv,
    research_queue,
    rows,
    save,
    seed,
    validate_database,
)


def test_seed_is_repeatable_and_has_no_fake_units(conn):
    candidates = json.loads(Path("data/seeds/buildings.json").read_text())
    assert seed(conn, candidates) == len(candidates)
    assert seed(conn, candidates) == 0
    assert not rows(conn, "units")
    assert not validate_database(conn)


def test_workbook_and_export_repeatability(conn):
    load_demo(conn)
    before = len(rows(conn, "audit"))
    for _ in range(2):
        book = load_workbook(io.BytesIO(workbook_bytes(conn)))
        assert len(book.sheetnames) == 12
        assert book["Buildings"].freeze_panes == "C4"
        assert book["Buildings"].auto_filter.ref
        assert book["Dashboard"]["B5"].value.startswith("READ-ONLY")
        assert book["Apartment Rankings"].max_row == 5
        assert all(c.data_type != "f" for sheet in book for row in sheet for c in row)
    assert len(rows(conn, "audit")) == before
    assert "'=HYPERLINK" in csv_bytes([{"note": '=HYPERLINK("bad")'}]).decode("utf-8-sig")


def test_adapter_reports_manual_requirement(conn):
    result = refresh_sources(conn)
    assert all(r.status == "manual_required" and r.changed == 0 for r in result)


def test_exported_entity_csv_round_trip(conn):
    from apartment_hunter.schemas import SCHEMAS

    load_demo(conn)
    for entity in ("buildings", "units", "contacts"):
        content = csv_bytes(rows(conn, entity), ["id", "revision", *SCHEMAS[entity].model_fields])
        assert import_csv(conn, entity, content.decode("utf-8-sig")) == 0


@pytest.mark.parametrize("populated", [False, True])
def test_dashboard_all_pages(tmp_path, monkeypatch, populated):
    path = tmp_path / "dashboard.db"
    monkeypatch.setenv("APARTMENT_DB", str(path))
    if populated:
        with database(path) as conn:
            load_demo(conn)
    app = AppTest.from_file(
        Path("src/apartment_hunter/dashboard.py").resolve(), default_timeout=30
    ).run()
    assert not app.exception
    for page in (
        "Review",
        "Buildings",
        "Property explorer",
        "Compare",
        "Map",
        "Contacts & tours",
        "Research queue",
        "Add / edit",
        "Import & export",
    ):
        app.sidebar.radio[0].set_value(page).run()
        assert not app.exception, f"Failed screen {page}: {app.exception}"


def test_dashboard_creates_building_and_unit(tmp_path, monkeypatch):
    path = tmp_path / "ui-write.db"
    monkeypatch.setenv("APARTMENT_DB", str(path))
    app = AppTest.from_file(
        Path("src/apartment_hunter/dashboard.py").resolve(), default_timeout=30
    ).run()
    app.sidebar.radio[0].set_value("Add / edit").run()
    app.text_input(key="buildings_0_name").set_value("UI example building")
    next(b for b in app.button if b.label == "Save record").click().run()
    assert not app.exception
    assert app.success
    with database(path) as conn:
        assert rows(conn, "buildings")[0]["name"] == "UI example building"
    app.selectbox(key="edit_entity").set_value("units").run()
    app.text_input(key="units_0_building_id").set_value("1")
    app.text_input(key="units_0_unit_number").set_value("2a")
    app.text_input(key="units_0_base_rent").set_value("3500")
    next(b for b in app.button if b.label == "Save record").click().run()
    assert not app.exception
    assert app.success
    with database(path) as conn:
        assert rows(conn, "units")[0]["unit_number"] == "2A"


def test_research_queue_puts_shortlist_first_and_skips_dropped(conn):
    for name, status in (("Plain", "considering"), ("Gone", "dropped"), ("Top", "shortlist")):
        building = save(conn, "buildings", {"name": name, "household_status": status})
        save(conn, "units", {"building_id": building, "unit_number": "1A"})
    queue = research_queue(conn)
    names = [task["building"] for task in queue if task["unit"] is None]
    assert names[0] == "Top" and "Plain" in names and "Gone" not in names
    gone = next(b["id"] for b in rows(conn, "buildings") if b["name"] == "Gone")
    assert all(task["building"] != gone for task in queue)
    questions = {task["question"] for task in queue}
    assert "Check on a map whether it is next to an elevated train" in questions
    assert "Confirm square footage" in questions
    assert "Check surroundings and unit views together" not in questions  # Cemetery relaxed.


def test_dashboard_review_page_marks_new_items_reviewed(tmp_path, monkeypatch):
    path = tmp_path / "review.db"
    monkeypatch.setenv("APARTMENT_DB", str(path))
    with database(path) as conn:
        load_demo(conn)
        save(conn, "buildings", {"name": "Seen before", "reviewed_on": today()})
        save(conn, "buildings", {"name": "Rejected", "household_status": "dropped"})
    app = AppTest.from_file(
        Path("src/apartment_hunter/dashboard.py").resolve(), default_timeout=30
    ).run()
    app.sidebar.radio[0].set_value("Review").run()
    assert not app.exception
    # New: the demo building and its two apartments. Reviewed: "Seen before" and "Rejected".
    assert [t.label for t in app.tabs] == ["New (3)", "Already reviewed (2)"]
    app.multiselect(key="reviewed").set_value([("buildings", 1), ("units", 2)])
    next(b for b in app.button if b.label == "Mark as reviewed today").click().run()
    assert not app.exception
    assert [t.label for t in app.tabs] == ["New (1)", "Already reviewed (4)"]
    assert app.success
    with database(path) as conn:
        reviewed = {(b["name"], b["reviewed_on"]) for b in rows(conn, "buildings")}
        assert ("SYNTHETIC Example House", today()) in reviewed
        assert [u["reviewed_on"] for u in rows(conn, "units")] == [None, today()]
